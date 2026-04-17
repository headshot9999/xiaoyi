from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

sections = doc.sections
for section in sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

def add_heading_styled(text, level=1, color=RGBColor(0, 51, 102)):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = color
        run.font.name = '微软雅黑'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return heading

def add_para(text, bold=False, italic=False, color=None, size=None, align=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = size
    if align:
        p.alignment = align
    p.paragraph_format.space_after = space_after
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run_bold = p.add_run(bold_prefix)
        run_bold.bold = True
        run_bold.font.name = '微软雅黑'
        run_bold.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(11)
    return p

def add_divider():
    p = doc.add_paragraph()
    run = p.add_run("─" * 50)
    run.font.color.rgb = RGBColor(200, 200, 200)
    run.font.size = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)


# ════════════════════════════════════════
# 封面
# ════════════════════════════════════════
for _ in range(5):
    doc.add_paragraph()

add_para("艾多美 AI 创意提案大赛", bold=True, size=Pt(28),
         color=RGBColor(0, 51, 102), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(24))

add_para("AI 时空穿梭机", bold=True, size=Pt(24),
         color=RGBColor(0, 102, 153), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(8))

add_para("当会员书写完人生剧本，AI 能为 TA 做什么？", bold=True, size=Pt(15),
         color=RGBColor(0, 102, 153), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(16))

add_para("——从一张雷达图，到一整个「可看见的未来」",
         italic=True, size=Pt(13),
         color=RGBColor(100, 100, 100), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(50))

add_para("提案人：_______________", size=Pt(12), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
add_para("部门：_______________", size=Pt(12), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
add_para("日期：2026年4月", size=Pt(12), align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ════════════════════════════════════════
# 目录
# ════════════════════════════════════════
add_heading_styled("目录", level=1)
toc_items = [
    "一、现状分析：会员书写完剧本之后，发生了什么？",
    "二、核心提案：AI 时空穿梭机——让剧本「活」起来",
    "三、AI 能力 1：梦想电影生成——把文字变成你主演的电影",
    "四、AI 能力 2：AI 智能教练——把剧本变成行动计划",
    "五、AI 能力 3：梦想进度追踪——把目标变成可量化的旅程",
    "六、AI 能力 4：AI 颁奖典礼——把成就变成人生高光时刻",
    "七、AI 能力 5：剧本社交——把个人梦想变成群体共振",
    "八、完整用户体验流程演示",
    "九、各维度 AI 赋能场景速查表",
    "十、技术实现方案",
    "十一、与艾多美业务的价值融合",
    "十二、实施路线图",
    "十三、总结",
]
for item in toc_items:
    add_para(item, size=Pt(11), space_after=Pt(3))

doc.add_page_break()

# ════════════════════════════════════════
# 一、现状分析
# ════════════════════════════════════════
add_heading_styled("一、现状分析：会员书写完剧本之后，发生了什么？", level=1)

add_heading_styled("1.1 现有流程", level=2, color=RGBColor(0, 102, 153))
add_para("目前，会员在手机端可以从以下维度中选择 8 个来书写自己的人生剧本：")

table = doc.add_table(rows=5, cols=2)
table.style = 'Light Shading Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['维度分类', '具体维度']
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h

dims = [
    ("家庭与关系", "妻子/丈夫、子女、父母、亲属、朋友"),
    ("个人发展", "事业、健康、教育/读书、个人成长、兴趣爱好、娱乐、旅行"),
    ("财务目标", "存款、房子、车"),
    ("社会责任", "志愿者服务、教育捐助、医疗捐助、环境保护、扶贫救济、灾害救助、特殊群体关爱"),
]
for i, (cat, detail) in enumerate(dims):
    table.rows[i+1].cells[0].text = cat
    table.rows[i+1].cells[1].text = detail

doc.add_paragraph()
add_para("会员选择8个维度 → 为每个维度书写剧本内容（文字描述 + 评分） → 系统生成一张雷达图。")

add_heading_styled("1.2 问题：然后呢？", level=2, color=RGBColor(0, 102, 153))
add_para("雷达图生成之后，会员的体验就结束了。", bold=True, color=RGBColor(200, 50, 50))
add_para("这里存在四个核心痛点：")

add_bullet("写完就忘，缺乏后续触达——剧本沉睡在App角落，没有持续激励机制",
           bold_prefix="痛点1  ")
add_bullet("雷达图太抽象，没有情感冲击——一张冰冷的数据图无法激发内心的驱动力",
           bold_prefix="痛点2  ")
add_bullet("目标没有拆解，不知道下一步做什么——写了「买房」两个字，然后呢？",
           bold_prefix="痛点3  ")
add_bullet("没有成就反馈，缺乏阶段性仪式感——达成了某个目标，没有任何庆祝",
           bold_prefix="痛点4  ")

add_heading_styled("1.3 机会", level=2, color=RGBColor(0, 102, 153))
add_para("会员已经书写了剧本内容——这是极其珍贵的「梦想数据」。"
         "每一段文字背后都是一个人对未来的渴望和想象。"
         "我们的AI不是要替代剧本，而是要让剧本书写完之后，真正产生价值。",
         bold=True, color=RGBColor(0, 51, 102))

doc.add_page_break()

# ════════════════════════════════════════
# 二、核心提案
# ════════════════════════════════════════
add_heading_styled("二、核心提案：AI 时空穿梭机——让剧本「活」起来", level=1)

add_para("我们提出「AI 时空穿梭机」方案：\n"
         "在会员书写完人生剧本的那一刻起，AI 接管后续体验，为每一段剧本内容提供五大AI能力加持。",
         size=Pt(12))

doc.add_paragraph()

add_para("现有流程 vs 升级后流程", bold=True, size=Pt(13), color=RGBColor(0, 51, 102))

table2 = doc.add_table(rows=3, cols=2)
table2.style = 'Light Shading Accent 1'
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
table2.rows[0].cells[0].text = '现状'
table2.rows[0].cells[1].text = 'AI 时空穿梭机升级后'
table2.rows[1].cells[0].text = '书写剧本 → 雷达图 → 结束'
table2.rows[1].cells[1].text = '书写剧本 → AI 即时响应 → 持续陪伴 → 目标达成'
table2.rows[2].cells[0].text = '一次性体验，看完就走'
table2.rows[2].cells[1].text = '剧本成为「活」的人生操作系统，每天都在运转'

doc.add_paragraph()

add_para("五大 AI 能力总览：", bold=True, size=Pt(12), color=RGBColor(0, 51, 102))

cap_table = doc.add_table(rows=6, cols=3)
cap_table.style = 'Light Shading Accent 1'
cap_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cap_headers = ['AI 能力', '一句话说明', '用户感受']
for i, h in enumerate(cap_headers):
    cap_table.rows[0].cells[i].text = h

caps = [
    ("1. 梦想电影生成", "把文字剧本变成你主演的电影", "「天呐，我看见未来的自己了！」"),
    ("2. AI 智能教练", "把剧本变成每天的行动计划", "「原来我每天该做这些事！」"),
    ("3. 梦想进度追踪", "把目标变成可量化的旅程", "「我离梦想还有 37% 的距离」"),
    ("4. AI 颁奖典礼", "把成就变成人生高光时刻", "「这就是我站在巅峰的样子！」"),
    ("5. 剧本社交", "把个人梦想变成群体共振", "「原来有人跟我一样的梦想！」"),
]
for i, (c1, c2, c3) in enumerate(caps):
    cap_table.rows[i+1].cells[0].text = c1
    cap_table.rows[i+1].cells[1].text = c2
    cap_table.rows[i+1].cells[2].text = c3

doc.add_paragraph()
add_para("以下逐一展开每个AI能力的详细设计。", italic=True, color=RGBColor(100, 100, 100))

doc.add_page_break()

# ════════════════════════════════════════
# 三、AI 能力 1：梦想电影
# ════════════════════════════════════════
add_heading_styled("三、AI 能力 1：梦想电影生成——把文字变成你主演的电影", level=1)

add_para("触发时机：会员书写完任意一个维度的剧本后，点击「生成我的梦想电影」按钮。",
         bold=True, color=RGBColor(0, 102, 153))

add_heading_styled("3.1 核心机制", level=2, color=RGBColor(0, 102, 153))
add_para("会员上传一张个人照片，AI 生成其「数字分身」。"
         "然后将剧本中的文字描述转化为一段 30-60 秒的 4K 电影级短片，"
         "数字分身在 AI 生成的未来场景中真实「出演」。")

add_heading_styled("3.2 各维度生成示例", level=2, color=RGBColor(0, 102, 153))

add_para("【家庭维度】妻子/丈夫", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「十周年纪念日，带她去马尔代夫，在水上别墅看日落。」\n"
         "→ AI 生成：傍晚的马尔代夫水上别墅，金色夕阳洒在碧蓝海面上。"
         "你和爱人的数字分身并肩坐在露台上，脚轻轻触碰海水，"
         "她靠在你肩上微笑。远处海豚跃出水面，背景音乐是轻柔的钢琴曲。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【家庭维度】子女", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「女儿考上理想的大学，送她去学校报到。」\n"
         "→ AI 生成：秋天的大学校园，银杏叶金黄。你帮女儿提着行李箱走在林荫道上，"
         "她兴奋地回头跟你说着什么。到了宿舍门口，你拍拍她的肩膀，"
         "她跑进去又折返回来给了你一个拥抱。你转身走的时候偷偷抹了下眼角。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【家庭维度】父母", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「给爸妈在老家盖一栋新房子，院子里种满妈妈喜欢的花。」\n"
         "→ AI 生成：一栋崭新的二层小楼，院子里鲜花盛开。"
         "妈妈系着围裙在花丛中浇水，爸爸坐在廊下的摇椅上喝茶。"
         "你开车回到门口，妈妈放下水壶迎上来。一家人坐在院子里吃晚饭，"
         "桌上摆满了菜，夕阳把所有人的影子拉得很长很长。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【财务维度】车", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「买一辆宝马X5，周末带全家自驾游。」\n"
         "→ AI 生成：一辆深蓝色宝马X5停在山间公路的观景台旁。"
         "你打开车门走出来伸个懒腰，远处是壮丽的山谷云海。"
         "孩子从后座蹦出来欢呼，爱人拿出保温杯递给你。"
         "一家人靠着车头，背后是无限风景。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【财务维度】房子", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「在城市里买一套大平层，有落地窗可以看江景。」\n"
         "→ AI 生成：清晨，阳光透过巨大的落地窗洒进客厅。"
         "你端着咖啡站在窗前，俯瞰蜿蜒的江面和城市天际线。"
         "身后是温馨的家——书架、绿植、孩子的玩具。"
         "镜头缓缓后退，这个家每一个角落都写满了幸福。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【个人发展】旅行", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「和闺蜜一起去冰岛看极光。」\n"
         "→ AI 生成：冰岛的冬夜，漫天极光如绿色的丝带在天空中飘舞。"
         "你和朋友的数字分身裹着羽绒服站在雪地里，仰头惊叹。"
         "她拉着你的手兴奋地跳起来，两个人在极光下笑得像孩子。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【社会责任】教育捐助", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「资助10个山区孩子完成学业。」\n"
         "→ AI 生成：一间明亮的山区教室，孩子们坐在崭新的课桌前认真听课。"
         "墙上贴着你捐赠的校名牌匾。画面切换到毕业典礼，"
         "那些孩子穿着学士服笑着向镜头挥手，手里举着「感谢有您」的牌子。"
         "你站在人群后面，安静地微笑着。",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("【个人发展】事业", bold=True, color=RGBColor(0, 51, 102))
add_para("剧本写：「三年内达成首席总监。」\n"
         "→ AI 触发「颁奖典礼」特殊模式（见第六章详解），生成一场完整的颁奖盛典电影。",
         italic=True, color=RGBColor(80, 80, 80))

add_divider()

add_heading_styled("3.3 关键设计：命运分岔路", level=2, color=RGBColor(0, 102, 153))
add_para("在任意维度的梦想电影生成后，用户可选择生成「双时间线对比版」——")
add_para("分屏同时展示「坚持之后的你」和「放弃之后的你」：")

add_para("例：健康维度", bold=True, color=RGBColor(0, 102, 0))
add_para("左屏（坚持）：你在晨光中慢跑，身材健康，精神饱满，陪孩子踢球。\n"
         "右屏（放弃）：你窝在沙发上刷手机，体检报告上亮了三个红灯。\n\n"
         "例：事业维度\n"
         "左屏（坚持）：你站在艾多美颁奖舞台上，灯光汇聚，全场掌声。\n"
         "右屏（放弃）：你坐在同一张办公桌前，看着同期加入的伙伴在朋友圈晒出晋升照片。",
         italic=True, color=RGBColor(80, 80, 80))

add_heading_styled("3.4 家庭合拍大片", level=2, color=RGBColor(0, 102, 153))
add_para("当多个家庭成员都有艾多美账号时，AI 可以自动将他们各自维度的剧本融合，"
         "生成一部「全家梦想合集」电影：爸爸的事业梦+妈妈的旅行梦+孩子的成长梦编织在一起。")

doc.add_page_break()

# ════════════════════════════════════════
# 四、AI 能力 2：AI 智能教练
# ════════════════════════════════════════
add_heading_styled("四、AI 能力 2：AI 智能教练——把剧本变成行动计划", level=1)

add_para("触发时机：会员书写完剧本后，AI 自动分析每个维度的目标，生成可执行的行动路径。",
         bold=True, color=RGBColor(0, 102, 153))

add_heading_styled("4.1 智能目标拆解", level=2, color=RGBColor(0, 102, 153))
add_para("会员写下的每一个梦想，AI 自动拆解为阶段目标 → 月度计划 → 每周/每日行动。")

add_para("示例：「三年后买一套 200 万的房子」", bold=True, color=RGBColor(0, 102, 0))
add_para("AI 拆解为：\n"
         "  · 目标首付：60万（30%）\n"
         "  · 当前存款：8万（差额52万）\n"
         "  · 需要月均储蓄：约14,500元/月\n"
         "  · 行动建议：\n"
         "    - 主业收入提升路径：当前级别→目标级别需要的PV和团队规模\n"
         "    - 开源路径：建议增加XX品类推广（利润率更高）\n"
         "    - 节流路径：AI分析消费习惯，给出合理的省钱建议",
         italic=True, color=RGBColor(80, 80, 80))

doc.add_paragraph()
add_para("示例：「资助10个山区孩子」", bold=True, color=RGBColor(0, 102, 0))
add_para("AI 拆解为：\n"
         "  · 每年资助费用估算：约2-3万元\n"
         "  · 对接渠道推荐：可合作的公益基金\n"
         "  · 建议路径：先从1个孩子开始，随事业发展逐步增加\n"
         "  · 时间线：第1年资助2个 → 第2年增至5个 → 第3年达成10个",
         italic=True, color=RGBColor(80, 80, 80))

add_heading_styled("4.2 AI 日常陪伴", level=2, color=RGBColor(0, 102, 153))
add_para("AI 教练不是一次性出报告，而是持续陪伴：")
add_bullet("晨间推送「今日行动卡」——根据当天日程，提醒该做的关键动作")
add_bullet("遇到低谷时推送之前生成的梦想电影——「记得你为什么出发」")
add_bullet("周末推送维度平衡提醒——「这周事业维度投入了40小时，但健康维度只走了3000步，"
           "建议周末安排一次运动」")

add_heading_styled("4.3 跨维度智能关联", level=2, color=RGBColor(0, 102, 153))
add_para("AI 不只是孤立看每个维度，而是发现维度之间的协同和冲突：")
add_bullet("协同发现：「你的事业目标和买房目标高度关联——"
           "达成钻石经销商后的佣金增长足以覆盖房贷首付差额」", bold_prefix="协同 → ")
add_bullet("冲突预警：「你同时规划了高强度事业冲刺和环球旅行，"
           "建议将旅行安排在达成阶段目标之后，作为奖励」", bold_prefix="冲突 → ")
add_bullet("捆绑建议：「你的健康目标是跑完马拉松，旅行目标是去东京——"
           "建议报名东京马拉松，一举两得」", bold_prefix="捆绑 → ")

doc.add_page_break()

# ════════════════════════════════════════
# 五、AI 能力 3：梦想进度追踪
# ════════════════════════════════════════
add_heading_styled("五、AI 能力 3：梦想进度追踪——把目标变成可量化的旅程", level=1)

add_para("触发时机：实时运行，持续追踪各维度进展。",
         bold=True, color=RGBColor(0, 102, 153))

add_heading_styled("5.1 从雷达图到「梦想星球」", level=2, color=RGBColor(0, 102, 153))
add_para("目前的雷达图是静态的、冰冷的。升级后：")

add_para("雷达图 → 进化为 3D「梦想星球」：", bold=True)
add_para("  · 8 个维度对应星球上的 8 个区域\n"
         "  · 每个区域随着进度推进会「生长」——房子维度从一块空地到地基到封顶\n"
         "  · 旅行维度每完成一个目的地，星球上就多一个地标\n"
         "  · 社会责任维度每完成一次捐助，天空中多一颗星星\n"
         "  · 星球有天气系统：你活跃的日子是晴天，松懈的日子阴天\n"
         "  · 用户可以在手机上旋转、缩放、探索自己的梦想星球")

add_heading_styled("5.2 AI 进度分析（取代纯数字）", level=2, color=RGBColor(0, 102, 153))
add_para("不再只是「完成 35%」这样的冰冷数字，AI 提供有温度的分析：")

add_para("示例：", bold=True, color=RGBColor(0, 102, 0))
add_para("「本月你的事业维度突飞猛进，销售额提升了10%，距离你买梦想之车'宝马X5'的目标"
         "又近了一步——按目前速度，预计还需18个月。继续保持！\n\n"
         "但我注意到你已经连续两周没有运动了。你的健康维度分数从上月的7分降到了5分。"
         "记得你剧本里写的吗？'我要活得健康，陪孩子一起长大。'今晚下楼走30分钟吧。」",
         italic=True, color=RGBColor(80, 80, 80))

add_heading_styled("5.3 维度平衡预警", level=2, color=RGBColor(0, 102, 153))
add_para("当某些维度长期停滞或失衡时，AI 主动提醒：")
add_bullet("「你的8个维度中，事业和财务发展很快（8分），但'朋友'维度已经3个月没有更新了。"
           "你的剧本里写过'定期和老同学聚会'——上次聚会是什么时候？」")
add_bullet("「你的'环境保护'维度目标是'每月参加一次环保活动'，但本季度只参加了1次。"
           "本周六社区有一场植树活动，要不要报名？」")

doc.add_page_break()

# ════════════════════════════════════════
# 六、AI 能力 4：AI 颁奖典礼
# ════════════════════════════════════════
add_heading_styled("六、AI 能力 4：AI 颁奖典礼——把成就变成人生高光时刻", level=1)

add_para("触发时机：①会员书写「事业」维度剧本后可预览颁奖预演；②真实达成晋升时自动生成。",
         bold=True, color=RGBColor(0, 102, 153))

add_heading_styled("6.1 为什么需要颁奖典礼？", level=2, color=RGBColor(0, 102, 153))
add_para("在艾多美的成长体系中，级别晋升是最重要的里程碑。但目前的庆祝方式主要停留在"
         "App通知和线下年会表彰。AI 颁奖典礼要做的是——\n\n"
         "在会员达成（或即将达成）晋升时，为 TA 生成一场 2-3 分钟的好莱坞品质专属颁奖盛典电影。"
         "TA 就是唯一的主角。")

add_heading_styled("6.2 颁奖典礼全流程", level=2, color=RGBColor(0, 102, 153))

scenes = [
    ("【开场：星光降临】",
     "航拍城市夜景，镜头推进到灯火通明的颁奖会场。"
     "巨幅LED屏上滚动着会员的名字和即将获得的头衔。交响乐序曲响起。"),
    ("【第一幕：王者座驾抵达】",
     "一辆豪车缓缓驶入（根据级别不同匹配不同座驾——首席总监对应劳斯莱斯幻影）。"
     "车门打开，会员的数字分身走出，两侧闪光灯和掌声齐鸣。"),
    ("【第二幕：红毯上的人生回顾】",
     "红毯两侧LED墙自动播放会员的奋斗时间线——加入的第一天、第一单、"
     "团队从1人到100人到1000人的关键节点。每走过一个里程碑，烟花绽放一次。"),
    ("【第三幕：会长亲迎】",
     "红毯尽头，朴韩吉会长的数字分身微笑等候。两人紧握双手、热情拥抱。"
     "会长拍着会员的肩膀，两人并肩登上舞台。特写：会长在耳边低语"
     "「你做到了，我一直相信你。」"),
    ("【第四幕：加冕时刻】",
     "会长亲手将荣誉徽章别在会员胸前。大屏展示成就数据。"
     "台下家人数字分身热泪盈眶、起立鼓掌。孩子兴奋地喊「爸爸/妈妈！」"),
    ("【第五幕：巅峰致辞】",
     "会员站在话筒前发表感言。AI 根据其个人经历自动生成感人致辞，"
     "呼应TA人生剧本中最初写下的那些梦想。"),
    ("【终幕：荣耀之夜】",
     "举杯同庆，漫天金箔飘落，烟花绽放。最后，会员走出会场，"
     "从口袋里掏出那张手写的人生剧本，贴在胸口，微笑着驶入夜色。"),
]

for title, desc in scenes:
    add_para(title, bold=True, color=RGBColor(0, 51, 102), size=Pt(11))
    add_para(desc, italic=True, color=RGBColor(80, 80, 80), space_after=Pt(8))

add_heading_styled("6.3 不同级别的差异化场景", level=2, color=RGBColor(0, 102, 153))

level_table = doc.add_table(rows=7, cols=4)
level_table.style = 'Light Shading Accent 1'
level_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(['达成级别', '座驾', '场景规格', '专属彩蛋']):
    level_table.rows[0].cells[i].text = h

levels = [
    ("经销商", "奔驰E级", "精致小型颁奖厅", "团队合影留念"),
    ("玫瑰经销商", "宝马7系", "中型颁奖会场", "玫瑰花雨特效"),
    ("红宝石经销商", "保时捷帕拉梅拉", "大型颁奖殿堂", "红宝石灯光秀"),
    ("钻石经销商", "兰博基尼/法拉利", "国际级颁奖盛典", "钻石烟花+无人机编队"),
    ("首席经销商", "宾利飞驰", "顶级殿堂盛典", "私人飞机开场画面"),
    ("首席总监", "劳斯莱斯幻影", "万人体育场超级盛典", "会长全程陪同\n直升机航拍"),
]
for i, (lv, car, scene, bonus) in enumerate(levels):
    level_table.rows[i+1].cells[0].text = lv
    level_table.rows[i+1].cells[1].text = car
    level_table.rows[i+1].cells[2].text = scene
    level_table.rows[i+1].cells[3].text = bonus

add_heading_styled("6.4 不止是事业——任意维度都有「庆祝时刻」", level=2, color=RGBColor(0, 102, 153))
add_para("颁奖典礼的逻辑不只适用于事业晋升，任何维度的重大目标达成都值得庆祝：", bold=True)

add_bullet("「房子」维度达成 → AI 生成你拿到新房钥匙、推开家门那一刻的微电影")
add_bullet("「车」维度达成 → AI 生成你坐进驾驶座、第一次驾车上路的兴奋画面")
add_bullet("「旅行」维度达成 → AI 生成旅行回忆影集，配以旅途中的高光镜头")
add_bullet("「教育捐助」达成 → AI 生成受助学生毕业、向你致谢的温暖画面")
add_bullet("「健康」维度达成 → AI 生成你冲过马拉松终点线的慢动作特写")
add_bullet("「子女」维度达成 → AI 生成孩子长大成才、回家团圆的感人画面")

doc.add_page_break()

# ════════════════════════════════════════
# 七、AI 能力 5：剧本社交
# ════════════════════════════════════════
add_heading_styled("七、AI 能力 5：剧本社交——把个人梦想变成群体共振", level=1)

add_para("触发时机：会员书写完剧本后可选择公开分享自己的梦想维度。",
         bold=True, color=RGBColor(0, 102, 153))

add_heading_styled("7.1 梦想匹配", level=2, color=RGBColor(0, 102, 153))
add_para("AI 分析所有会员的剧本内容，发现「同频」的人：")
add_bullet("「有 23 位会员和你一样选择了'环境保护'维度，要加入环保梦想圈？」")
add_bullet("「你们团队有 5 个人都写了'明年去日本旅行'——要不要组团？」")
add_bullet("「你和李姐都在'子女教育'维度写了'让孩子学钢琴'，她的孩子已经学了两年，可以请教」")

add_heading_styled("7.2 梦想电影互看", level=2, color=RGBColor(0, 102, 153))
add_para("在用户授权的前提下，可以观看其他会员的梦想电影片段：")
add_bullet("团队会议上播放成员们的梦想电影合辑——这是最强的团队激励")
add_bullet("新人加入后看到前辈的梦想电影和实现过程——快速建立信任和目标感")
add_bullet("在社交媒体分享自己的梦想电影——天然的裂变传播素材")

add_heading_styled("7.3 时间胶囊", level=2, color=RGBColor(0, 102, 153))
add_para("每位会员可以在书写剧本时，同步录制一段「给未来自己的话」。"
         "AI 加密保存，到达设定日期后自动解锁，并将当时的誓言与实际成就进行对比，"
         "生成一段「预言成真」回顾视频。这样的内容一旦分享到社区，感染力极强。")

doc.add_page_break()

# ════════════════════════════════════════
# 八、完整用户体验流程演示
# ════════════════════════════════════════
add_heading_styled("八、完整用户体验流程演示", level=1)

add_para("以会员「小美」为例，完整展示从书写剧本到AI赋能的全流程：",
         italic=True, color=RGBColor(100, 100, 100))

steps = [
    ("Step 1  打开App，选择 8 个维度",
     "小美选择了：妻子（改为丈夫视角——她已婚）、子女、父母、朋友、事业、健康、房子、教育捐助。"),
    ("Step 2  逐一书写每个维度的剧本",
     "事业维度：「三年内达成钻石经销商」\n"
     "房子维度：「在杭州买一套西湖附近的房子」\n"
     "父母维度：「每年带爸妈出国旅行一次」\n"
     "教育捐助：「资助5个山区孩子上学」\n"
     "……"),
    ("Step 3  AI 即时响应——雷达图升级",
     "书写完成后，不再只有一张雷达图：\n"
     "  ✦ 雷达图保留（作为基础数据面板）\n"
     "  ✦ 新增「生成我的梦想电影」按钮 → 点击后，30秒内先出预览海报，5分钟内出完整短片\n"
     "  ✦ 新增「查看我的行动计划」→ AI 教练自动生成8个维度的拆解目标\n"
     "  ✦ 新增「探索我的梦想星球」→ 进入3D互动界面"),
    ("Step 4  小美点击「生成梦想电影」",
     "AI 根据她写的8个维度剧本，自动生成一段 2 分钟的「人生剧本电影」：\n"
     "  片段1：她带着爸妈在巴黎铁塔下的笑脸\n"
     "  片段2：她站在杭州新房的落地窗前看西湖\n"
     "  片段3：她在艾多美颁奖舞台上，灯光汇聚\n"
     "  片段4：山区教室里孩子们举着「感谢有您」的牌子\n"
     "  ……每个片段中都是她自己的数字分身「出演」"),
    ("Step 5  日常陪伴",
     "接下来的每一天：\n"
     "  · 早上收到AI教练推送：「今天的关键行动：联系2位新客户+30分钟运动」\n"
     "  · 周末收到维度平衡提醒：「本周事业进步很大，但健康维度落后了」\n"
     "  · 每月收到进度报告：「你距离钻石经销商还有38%的旅程」"),
    ("Step 6  达成里程碑——AI 颁奖",
     "一年后，小美达成了玫瑰经销商！\n"
     "  · AI 自动生成专属颁奖典礼电影（宝马7系+玫瑰花雨+会长握手）\n"
     "  · 她激动地分享到朋友圈和团队群\n"
     "  · 团队新人看到后纷纷打开自己的人生剧本，写下新的目标\n"
     "  · 裂变开始……"),
]

for title, desc in steps:
    add_para(title, bold=True, color=RGBColor(0, 51, 102), size=Pt(12))
    add_para(desc, color=RGBColor(80, 80, 80), space_after=Pt(12))

doc.add_page_break()

# ════════════════════════════════════════
# 九、各维度 AI 赋能场景速查表
# ════════════════════════════════════════
add_heading_styled("九、各维度 AI 赋能场景速查表", level=1)

add_para("下表展示了每个维度对应的 AI 能力应用场景，帮助评审快速理解AI赋能的广度和深度。",
         italic=True, color=RGBColor(100, 100, 100))

big_table = doc.add_table(rows=13, cols=4)
big_table.style = 'Light Shading Accent 1'
big_table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(['维度', '梦想电影示例', 'AI教练示例', '达成庆祝']):
    big_table.rows[0].cells[i].text = h

dim_rows = [
    ("妻子/丈夫", "结婚纪念日海外旅行", "每周安排一次约会之夜", "纪念日微电影"),
    ("子女", "孩子毕业典礼/成才画面", "制定亲子时光计划", "成长里程碑短片"),
    ("父母", "给父母盖新房/出国旅行", "每月存孝心基金", "全家福感恩电影"),
    ("朋友", "老友重聚/一起旅行", "定期聚会提醒", "友情纪念册"),
    ("事业", "站在颁奖舞台上的你", "PV目标拆解+日常行动", "AI颁奖典礼盛典"),
    ("健康", "马拉松终点/健康体检全绿", "运动+饮食每日打卡", "冲过终点线慢镜头"),
    ("房子", "推开新家门的那一刻", "首付储蓄计划+房贷测算", "新家入住微电影"),
    ("车", "第一次驾车上路/自驾游", "购车储蓄计划", "驾车出发公路电影"),
    ("旅行", "在目的地的美好画面", "旅行基金+请假规划", "旅行回忆影集"),
    ("教育捐助", "受助学生毕业向你致谢", "每月捐助计划+对接渠道", "公益成就纪录片"),
    ("志愿者服务", "你在公益现场帮助他人", "每月志愿时间规划", "志愿者荣誉影片"),
    ("环境保护", "你种下的树长成大树", "环保行动打卡计划", "一棵树到一片林的延时摄影"),
]
for i, (d, movie, coach, cele) in enumerate(dim_rows):
    big_table.rows[i+1].cells[0].text = d
    big_table.rows[i+1].cells[1].text = movie
    big_table.rows[i+1].cells[2].text = coach
    big_table.rows[i+1].cells[3].text = cele

doc.add_page_break()

# ════════════════════════════════════════
# 十、技术实现方案
# ════════════════════════════════════════
add_heading_styled("十、技术实现方案", level=1)

add_para("核心技术栈：", bold=True)
tech_items = [
    ("AI 视频生成：", "基于 Sora / Kling / Veo 2.0，支持文字→分镜→视频自动化生产，4K/60fps"),
    ("数字分身：", "单张照片 3D 面部重建 + LoRA 微调，支持表情控制和年龄渐变"),
    ("AI 教练引擎：", "大语言模型（GPT-4o / Claude）驱动目标拆解、行动推荐和进度分析"),
    ("3D 梦想星球：", "Three.js / Unity WebGL 渲染，手机浏览器端流畅运行"),
    ("AI 配乐：", "Suno / Udio 根据场景情绪实时生成原创配乐"),
    ("数据集成：", "对接艾多美业务系统 API，实时获取会员业绩、级别等数据"),
    ("隐私安全：", "端到端加密，面部特征不存储原始照片，所有数据加密传输"),
]
for prefix, desc in tech_items:
    add_bullet(desc, bold_prefix=prefix)

doc.add_paragraph()
add_para("关键体验指标：", bold=True)
add_bullet("书写完剧本 → 10秒内生成AI海报预览 → 5分钟内生成30秒短片")
add_bullet("颁奖典礼完整电影（2-3分钟）→ 30分钟内生成")
add_bullet("AI教练目标拆解 → 书写完成后即时呈现")
add_bullet("梦想星球 → 实时更新，手机端流畅交互")

doc.add_page_break()

# ════════════════════════════════════════
# 十一、与艾多美业务的价值融合
# ════════════════════════════════════════
add_heading_styled("十一、与艾多美业务的价值融合", level=1)

add_heading_styled("11.1 对会员", level=2, color=RGBColor(0, 102, 153))
add_bullet("人生剧本不再是一次性作业，而是每天在运转的「人生操作系统」",
           bold_prefix="剧本活化：")
add_bullet("AI教练将宏大梦想拆解为每天可以做到的小事，降低行动门槛",
           bold_prefix="行动落地：")
add_bullet("梦想电影、颁奖典礼视频成为个人最独特的社交内容",
           bold_prefix="个人品牌：")
add_bullet("任何维度的进步都被看见、被追踪、被庆祝，持续获得成就感",
           bold_prefix="持续激励：")

add_heading_styled("11.2 对团队", level=2, color=RGBColor(0, 102, 153))
add_bullet("播放成员的梦想电影是最强的团队激励手段",
           bold_prefix="激励工具：")
add_bullet("新人书写完剧本后立刻看到梦想电影，归属感瞬间建立",
           bold_prefix="新人留存：")
add_bullet("AI教练让每位成员的行动更有方向，团队整体效率提升",
           bold_prefix="效率提升：")

add_heading_styled("11.3 对公司", level=2, color=RGBColor(0, 102, 153))
add_bullet("全球直销行业首家AI驱动的「人生剧本操作系统」",
           bold_prefix="行业首创：")
add_bullet("用户自发分享的梦想电影和颁奖视频自带传播属性",
           bold_prefix="自传播：")
add_bullet("将「书写人生剧本」从企业文化口号升级为可体验、可衡量的AI产品",
           bold_prefix="品牌升级：")
add_bullet("海量的梦想维度数据为公司洞察会员需求提供支撑",
           bold_prefix="数据价值：")

doc.add_page_break()

# ════════════════════════════════════════
# 十二、实施路线图
# ════════════════════════════════════════
add_heading_styled("十二、实施路线图", level=1)

phases = [
    ("第一阶段：MVP 验证",
     ["在现有雷达图页面下方新增「生成我的梦想电影」按钮",
      "基础版梦想电影：文字 → AI海报 → 15秒短视频（先跑通核心链路）",
      "基础版AI教练：书写完剧本后自动生成8个维度的目标拆解",
      "100名种子用户内测"]),
    ("第二阶段：核心能力完善",
     ["数字分身系统上线（上传照片→个人分身出演）",
      "完整版梦想电影（30-60秒，多场景支持）",
      "AI颁奖典礼完整6幕电影",
      "梦想星球3D界面Beta版",
      "AI教练日常推送系统",
      "1000名用户公测"]),
    ("第三阶段：全面上线 + 社交生态",
     ["正式对全量用户开放",
      "剧本社交功能（梦想匹配、互看、时间胶囊）",
      "家庭合拍大片功能",
      "命运分岔路双时间线对比",
      "与线下颁奖活动联动",
      "多语言支持（覆盖艾多美全球市场）"]),
]

for title, items in phases:
    add_para(f"▎{title}", bold=True, color=RGBColor(0, 102, 153), size=Pt(12))
    for item in items:
        add_bullet(item)
    doc.add_paragraph()

doc.add_page_break()

# ════════════════════════════════════════
# 十三、总结
# ════════════════════════════════════════
add_heading_styled("十三、总结", level=1)

add_para("会员打开手机，选择8个维度，写下了关于未来的文字。\n"
         "今天，这些文字变成一张雷达图，然后——就没有然后了。",
         size=Pt(12))
doc.add_paragraph()

add_para("AI 时空穿梭机要做的，就是让「然后」发生：", bold=True, size=Pt(13),
         color=RGBColor(0, 51, 102))
doc.add_paragraph()

afters = [
    "写完「父母」维度 → 看见自己带爸妈在巴黎铁塔下的笑脸",
    "写完「房子」维度 → 看见自己推开新家大门那一刻的喜悦",
    "写完「事业」维度 → 看见自己开着劳斯莱斯驶上红毯、会长为你加冕",
    "写完「教育捐助」→ 看见受助的孩子穿上学士服，向你挥手致谢",
    "写完「健康」维度 → 看见自己冲过马拉松终点线的热泪和笑容",
    "写完「子女」维度 → 看见孩子长大成才、回家团圆的温馨画面",
]
for line in afters:
    add_para(f"  ✦  {line}", color=RGBColor(0, 51, 102), size=Pt(11), space_after=Pt(8))

doc.add_paragraph()

add_para("不只是看见——\n"
         "AI 教练帮你把每个梦想拆解成今天就能做的事。\n"
         "梦想星球帮你追踪每一步的进度。\n"
         "每一次达成，AI 都为你举办一场专属的颁奖典礼。\n"
         "你的剧本不再沉睡，它每天都在运行。",
         bold=True, size=Pt(12), color=RGBColor(0, 51, 102),
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(16))

add_para("让每一位艾多美人书写的人生剧本，\n"
         "不止于一张雷达图，\n"
         "而成为一整个「可看见、可行动、可庆祝」的未来。",
         bold=True, size=Pt(14), color=RGBColor(200, 50, 50),
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(30))

add_divider()
add_para("—— 感谢评审老师的审阅 ——", align=WD_ALIGN_PARAGRAPH.CENTER,
         color=RGBColor(150, 150, 150), italic=True)

# ════════════════════════════════════════
# 保存
# ════════════════════════════════════════
output_path = "/workspace/艾多美AI创意提案_AI时空穿梭机.docx"
doc.save(output_path)
print(f"文档已保存至: {output_path}")
